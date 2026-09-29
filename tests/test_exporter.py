"""Unit test for exporting to Excel and CSV."""
import os
import openpyxl
import pandas as pd
from core.models import Lead, LeadPriority, WebsiteStatus
from core.exporter import export_to_excel, export_to_csv


def test_export_excel(tmp_path):
    leads = [
        Lead(
            name="Горячий Клиент (Без сайта)",
            source="Яндекс Карты",
            category="Стоматология",
            city="Москва",
            address="ул. Арбат, 10",
            phones=["+7 (999) 111-22-33"],
            website=None,
            rating=4.9,
            reviews_count=45,
            map_url="https://yandex.ru/maps/org/123",
        ),
        Lead(
            name="Теплый Клиент (Таплинк)",
            source="2ГИС",
            category="Салон красоты",
            city="Москва",
            address="ул. Ленина, 5",
            phones=["+7 (988) 222-33-44"],
            website="https://taplink.cc/beauty",
            rating=4.8,
            reviews_count=20,
            map_url="https://2gis.ru/firm/456",
        ),
        Lead(
            name="Обычный Клиент (Есть сайт)",
            source="2ГИС",
            category="Автосервис",
            city="Москва",
            address="ш. Энтузиастов, 12",
            phones=["+7 (495) 777-88-99"],
            website="https://autoservice-expert.ru",
            rating=4.5,
            reviews_count=100,
        ),
    ]

    excel_file = str(tmp_path / "test_leads.xlsx")
    saved_path = export_to_excel(leads, excel_file)
    assert os.path.exists(saved_path)

    # Validate Excel content
    wb = openpyxl.load_workbook(saved_path)
    ws = wb.active
    assert ws.title == "Лиды для сайтов"
    assert ws.cell(row=1, column=1).value == "Приоритет лида"
    assert ws.cell(row=2, column=1).value == LeadPriority.HIGH.value
    assert ws.cell(row=2, column=3).value == "Горячий Клиент (Без сайта)"
    assert ws.cell(row=2, column=4).value == "📱 Моб."
    assert ws.cell(row=2, column=5).value == "+7 (999) 111-22-33"

    # CSV Test
    csv_file = str(tmp_path / "test_leads.csv")
    csv_saved = export_to_csv(leads, csv_file)
    assert os.path.exists(csv_saved)

    df = pd.read_csv(csv_saved, sep=";", encoding="utf-8-sig")
    assert len(df) == 3
    assert df.iloc[0]["Название"] == "Горячий Клиент (Без сайта)"
    print("Export tests passed successfully!")


if __name__ == "__main__":
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as td:
        test_export_excel(Path(td))
