"""Export leads to styled Excel and CSV formats."""
import os
from typing import List
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Lead, LeadPriority, WebsiteStatus, FreelanceOrder


def export_to_excel(leads: List[Lead], output_path: str = None) -> str:
    """Exports a list of leads to a professionally styled Excel workbook.

    Returns the absolute path to the generated file.
    """
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.abspath(f"leads_{timestamp}.xlsx")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды для сайтов"
    ws.views.sheetView[0].showGridLines = True

    # Column headers
    headers = [
        ("Приоритет лида", 24),
        ("Статус сайта", 22),
        ("Название компании", 32),
        ("Тип связи", 16),
        ("Телефоны", 26),
        ("WhatsApp", 24),
        ("Telegram", 24),
        ("Сайт компании", 28),
        ("Аудит сайта / Реклама", 28),
        ("AI-Питч (WhatsApp/TG)", 55),
        ("Скрипт звонка (30 сек)", 50),
        ("Ниша / Рубрика", 24),
        ("Город", 18),
        ("Адрес", 38),
        ("Кол-во отзывов", 16),
        ("Рейтинг", 12),
        ("Соцсети", 30),
        ("Режим работы", 22),
        ("Источник", 16),
        ("Ссылка на картах", 30),
    ]

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Priority fills
    fill_high = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")     # Soft green
    font_high = Font(name="Calibri", size=10, bold=True, color="276A3C")

    fill_medium = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # Soft yellow
    font_medium = Font(name="Calibri", size=10, bold=True, color="8A6D3B")

    fill_low = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")      # Soft gray
    font_low = Font(name="Calibri", size=10, color="595959")

    regular_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    mobile_font = Font(name="Calibri", size=10, bold=True, color="107C41")

    # Write headers
    for col_idx, (header_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28

    # Write data rows
    for row_idx, lead in enumerate(leads, start=2):
        wa_url = lead.whatsapp_url or "—"
        tg_url = lead.telegram_url or "—"
        row_data = [
            lead.lead_priority.value,
            lead.website_status.value,
            lead.name,
            lead.phone_type_badge,
            lead.phones_str,
            wa_url,
            tg_url,
            lead.website or "—",
            lead.site_audit_summary or "—",
            lead.ai_pitch or "—",
            lead.call_script or "—",
            lead.category,
            lead.city,
            lead.address,
            lead.reviews_count if lead.reviews_count is not None else 0,
            lead.rating if lead.rating is not None else "—",
            lead.socials_str or "—",
            lead.working_hours or "—",
            lead.source,
            lead.map_url or "—",
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.font = regular_font
            cell.alignment = Alignment(vertical="center")

            # Style Priority and Website Status columns
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if lead.lead_priority == LeadPriority.HIGH:
                    cell.fill = fill_high
                    cell.font = font_high
                elif lead.lead_priority == LeadPriority.MEDIUM:
                    cell.fill = fill_medium
                    cell.font = font_medium
                else:
                    cell.fill = fill_low
                    cell.font = font_low

            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            elif col_idx == 4:  # Phone type
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if lead.is_mobile:
                    cell.font = mobile_font

            elif col_idx == 5:  # Phones
                cell.alignment = Alignment(horizontal="left", vertical="center")

            elif col_idx == 6 and wa_url != "—":  # WhatsApp
                cell.font = link_font
                if wa_url.startswith("http"):
                    cell.hyperlink = wa_url

            elif col_idx == 7 and tg_url != "—":  # Telegram
                cell.font = link_font
                if tg_url.startswith("http"):
                    cell.hyperlink = tg_url

            elif col_idx == 8 and lead.website and lead.website != "—":  # Website
                cell.font = link_font
                if lead.website.startswith("http"):
                    cell.hyperlink = lead.website

            elif col_idx in (10, 11):  # AI Pitch & Call Script
                cell.alignment = Alignment(vertical="center", wrap_text=True)

            elif col_idx in (15, 16):  # Reviews & Rating
                cell.alignment = Alignment(horizontal="center", vertical="center")

            elif col_idx == 20 and lead.map_url and lead.map_url != "—":  # Map URL
                cell.font = link_font
                cell.hyperlink = lead.map_url

        ws.row_dimensions[row_idx].height = 24

    # Enable filter and freeze top row
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(output_path)
    return output_path


def export_to_csv(leads: List[Lead], output_path: str = None) -> str:
    """Exports leads to CSV with utf-8-sig encoding for perfect Russian Excel display."""
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.abspath(f"leads_{timestamp}.csv")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    rows = []
    for lead in leads:
        rows.append({
            "Приоритет лида": lead.lead_priority.value,
            "Статус сайта": lead.website_status.value,
            "Название": lead.name,
            "Тип связи": lead.phone_type_badge,
            "Телефоны": lead.phones_str,
            "WhatsApp": lead.whatsapp_url or "",
            "Telegram": lead.telegram_url or "",
            "Сайт": lead.website or "",
            "Аудит сайта": lead.site_audit_summary or "",
            "AI-Питч (WhatsApp/TG)": lead.ai_pitch or "",
            "Скрипт звонка": lead.call_script or "",
            "Ниша / Рубрика": lead.category,
            "Город": lead.city,
            "Адрес": lead.address,
            "Отзывы": lead.reviews_count if lead.reviews_count is not None else 0,
            "Рейтинг": lead.rating if lead.rating is not None else "",
            "Соцсети": lead.socials_str,
            "Время работы": lead.working_hours or "",
            "Источник": lead.source,
            "Ссылка на картах": lead.map_url or "",
        })

    headers = [
        "Приоритет лида", "Статус сайта", "Название", "Тип связи", "Телефоны",
        "WhatsApp", "Telegram", "Сайт", "Аудит сайта", "AI-Питч (WhatsApp/TG)",
        "Скрипт звонка", "Ниша / Рубрика", "Город", "Адрес", "Отзывы", "Рейтинг",
        "Соцсети", "Время работы", "Источник", "Ссылка на картах"
    ]
    with open(output_path, mode="w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return output_path


def export_freelance_to_excel(orders: List[FreelanceOrder], output_path: str = None) -> str:
    """Exports a list of freelance orders to a styled Excel workbook."""
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.abspath(f"freelance_orders_{timestamp}.xlsx")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы на разработку сайтов"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        ("Биржа", 18),
        ("Название заказа", 45),
        ("Бюджет / Оплата", 22),
        ("Ссылка на проект", 32),
        ("Описание задания", 60),
        ("Откликов", 16),
        ("Дата публикации", 22),
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    price_font = Font(name="Calibri", size=10, bold=True, color="15803D")
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")

    price_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    platform_fills = {
        "Onliner": PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid"),
        "Kufar": PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid"),
        "Telegram": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        "Kwork": PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),
        "FL.ru": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
        "Freelancehunt": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
        "Weblancer": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
    }

    for col_idx, (header_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28

    for row_idx, order in enumerate(orders, start=2):
        row_data = [
            order.platform,
            order.title,
            order.price or "По договоренности",
            order.url,
            order.description or "—",
            order.proposals_count or "—",
            order.date_posted or "—",
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.font = regular_font
            cell.alignment = Alignment(vertical="center")

            # Platform
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
                if order.platform in platform_fills:
                    cell.fill = platform_fills[order.platform]

            # Title
            elif col_idx == 2:
                cell.font = bold_font

            # Price
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = price_font
                cell.fill = price_fill

            # URL
            elif col_idx == 4:
                cell.value = "Открыть заказ ↗"
                cell.hyperlink = order.url
                cell.font = link_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Description
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            # Proposals & Date
            elif col_idx in (6, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_idx].height = 24

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path


def export_freelance_to_csv(orders: List[FreelanceOrder], output_path: str = None) -> str:
    """Exports freelance orders to CSV with utf-8-sig encoding."""
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.abspath(f"freelance_orders_{timestamp}.csv")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    rows = []
    for o in orders:
        rows.append({
            "Биржа": o.platform,
            "Название": o.title,
            "Бюджет": o.price,
            "Ссылка": o.url,
            "Описание": o.description,
            "Откликов": o.proposals_count,
            "Дата публикации": o.date_posted,
        })

    headers = ["Биржа", "Название", "Бюджет", "Ссылка", "Описание", "Откликов", "Дата публикации"]
    with open(output_path, mode="w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return output_path

